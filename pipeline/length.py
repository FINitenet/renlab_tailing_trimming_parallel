import logging
import os
from collections import Counter, OrderedDict
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/renlab-tailing-matplotlib")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pysam
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

from .common import run_parallel


def _build_end_matrix(meta_file):
    result = {}
    with open(meta_file) as handle:
        for line in handle:
            fields = line.rstrip().split("\t")
            parts = fields[0].split("-")
            hairpin = "-".join(parts[:2])
            result[f"{hairpin}-{fields[2]}"] = int(fields[2]) + len(fields[1]) - 1
    return result


def _read_existing_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        total = int(workbook["total reads mapped to miRNA"]["B1"].value or 0)
        distributions = []
        for sheet_name in ("distribution of whole reads", "distribution of 5GMC"):
            sheet = workbook[sheet_name]
            counts = OrderedDict()
            for length, count, _ in sheet.iter_rows(min_row=2, values_only=True):
                if length is not None and count is not None:
                    counts[int(length)] = int(count)
            distributions.append(counts)
        return total, distributions[0], distributions[1]
    finally:
        workbook.close()


def _add_excel_chart(summary_sheet, data_sheet, title, anchor):
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.x_axis.title = "Length (nt)"
    chart.y_axis.title = "Percentage of mapped miRNA reads"
    chart.height = 8
    chart.width = 14
    data = Reference(data_sheet, min_col=3, min_row=1, max_row=data_sheet.max_row)
    categories = Reference(data_sheet, min_col=1, min_row=2, max_row=data_sheet.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.legend = None
    summary_sheet.add_chart(chart, anchor)


def _write_workbook(path, total, read_lengths, gmc_lengths):
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "total reads mapped to miRNA"
    summary_sheet.append(("total reads mapped to miRNA:", total))
    summary_sheet["A3"] = "The charts below use the data stored in the other two worksheets."
    read_sheet = workbook.create_sheet("distribution of whole reads")
    gmc_sheet = workbook.create_sheet("distribution of 5GMC")
    for sheet, counts in ((read_sheet, read_lengths), (gmc_sheet, gmc_lengths)):
        sheet.append(("Length", "Count", "Percentage"))
        for length, count in OrderedDict(sorted(counts.items())).items():
            sheet.append((length, count, count / total if total else 0))
        for cell in sheet["C"][1:]:
            cell.number_format = "0.00%"
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 12
        sheet.column_dimensions["B"].width = 16
        sheet.column_dimensions["C"].width = 16
    _add_excel_chart(summary_sheet, read_sheet, "Whole-read length distribution", "A5")
    _add_excel_chart(summary_sheet, gmc_sheet, "5GMC length distribution", "A21")
    summary_sheet.column_dimensions["A"].width = 54
    summary_sheet.column_dimensions["B"].width = 18
    workbook.save(path)


def _plot_distributions(path, sample, total, read_lengths, gmc_lengths):
    path.parent.mkdir(parents=True, exist_ok=True)
    figure, axes = plt.subplots(1, 2, figsize=(11, 4.5), sharey=True)
    panels = (
        (axes[0], read_lengths, "Whole reads", "#4c78a8"),
        (axes[1], gmc_lengths, "5GMC", "#9b4b8f"),
    )
    for axis, counts, title, color in panels:
        lengths = sorted(counts)
        percentages = [(counts[length] / total * 100) if total else 0 for length in lengths]
        axis.bar(lengths, percentages, width=0.8, color=color, edgecolor="white", linewidth=0.4)
        axis.set_title(title, fontweight="bold")
        axis.set_xlabel("Length (nt)")
        axis.set_xticks(lengths)
        axis.tick_params(axis="x", labelrotation=45)
        axis.grid(axis="y", color="#d0d0d0", linewidth=0.5, alpha=0.8)
        axis.set_axisbelow(True)
    axes[0].set_ylabel("Percentage of mapped miRNA reads (%)")
    figure.suptitle(f"{sample} — read length distributions", fontweight="bold")
    figure.tight_layout(rect=(0, 0, 1, 0.94))
    figure.savefig(path, bbox_inches="tight")
    plt.close(figure)


def run_sample(item, config):
    sample, _ = item
    sam_file = Path(config["remapping_dir"]) / sample / "merged-alignment-sorted.sam"
    output_file = Path(config["length_dir"]) / f"{sample}_len_dist.xlsx"
    plot_file = Path(config["length_dir"]) / "plots" / f"{sample}_length_distribution.pdf"
    if config["resume"] and output_file.is_file() and plot_file.is_file():
        logging.info("[%s] length distribution already complete; skipping", sample)
        return
    if config["resume"] and output_file.is_file() and not plot_file.is_file():
        logging.info("[%s] length table exists; generating missing charts", sample)
        total, read_lengths, gmc_lengths = _read_existing_workbook(output_file)
        _write_workbook(output_file, total, read_lengths, gmc_lengths)
        _plot_distributions(plot_file, sample, total, read_lengths, gmc_lengths)
        logging.info("[%s] length charts complete", sample)
        return
    if config["dry_run"]:
        logging.info(
            "Would calculate length distributions: %s -> %s and %s",
            sam_file,
            output_file,
            plot_file,
        )
        return
    if not sam_file.is_file():
        raise FileNotFoundError(f"Remapping SAM not found for {sample}: {sam_file}")

    ends = _build_end_matrix(config["meta_file"])
    read_lengths = Counter()
    gmc_lengths = Counter()
    seen = set()
    total = 0
    with pysam.AlignmentFile(str(sam_file), "r") as alignments:
        for record in alignments:
            mirna_id = f"{record.reference_name}-{record.reference_start + 1}"
            name_parts = record.query_name.split("-")
            query_id = "-".join(name_parts[:2])
            query_cut = int(record.query_name.split("--")[1])
            raw_length = len(name_parts[2])
            end_5gmc = record.reference_start + raw_length - query_cut
            length_5gmc = raw_length - query_cut
            if mirna_id in ends and query_id not in seen:
                if query_cut == 0 and end_5gmc > ends[mirna_id]:
                    length_5gmc = raw_length
                read_lengths[raw_length] += 1
                gmc_lengths[length_5gmc] += 1
                total += 1
            seen.add(query_id)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    _write_workbook(output_file, total, read_lengths, gmc_lengths)
    _plot_distributions(plot_file, sample, total, read_lengths, gmc_lengths)
    logging.info("[%s] length distribution complete", sample)


def run(samples, config, jobs):
    run_parallel(run_sample, samples, config, jobs, "Length analysis")
